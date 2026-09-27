"""Identify exactly which shapes overflow, with their geometry."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.util import Emu

DECK = Path(__file__).resolve().parents[1] / "docs" / "presentation" / "Sistem-Informasi-Data-Siswa-Presentation.pptx"
EMU = 914400

prs = Presentation(str(DECK))
W, H = prs.slide_width, prs.slide_height
print(f"slide height: {H/EMU:.3f} in")

for i, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if shape.top is None or shape.height is None:
            continue
        bottom = shape.top + shape.height
        if bottom > H + Emu(int(0.05 * EMU)):
            txt = ""
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip().replace("\n", " ")[:48]
            kind = shape.__class__.__name__
            print(f"  slide {i}: {kind:16} top={shape.top/EMU:.2f} "
                  f"h={shape.height/EMU:.2f} bottom={bottom/EMU:.2f}  '{txt}'")
